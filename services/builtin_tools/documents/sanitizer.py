"""PDF·DOCX·XLSX의 정책 대상 메타데이터를 제거한다."""

from __future__ import annotations

from io import BytesIO
from zipfile import BadZipFile, ZIP_DEFLATED, ZipFile

from docx import Document
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pypdf import PdfReader, PdfWriter
from pypdf.errors import PdfReadError

from services.builtin_tools.common.errors import BuiltinToolError
from services.builtin_tools.common.limits import MAX_FILE_BYTES, MAX_OUTPUT_BYTES
from services.builtin_tools.documents.reader import DOCX_MIME, PDF_MIME, XLSX_MIME
from services.builtin_tools.documents.package_safety import validate_ooxml_package
from services.builtin_tools.documents.metadata import sensitive_metadata


def _sanitize_pdf(data: bytes) -> bytes:
    try:
        reader = PdfReader(BytesIO(data))
    except PdfReadError as exc:
        raise BuiltinToolError("INVALID_PDF", "손상되었거나 지원하지 않는 PDF입니다.") from exc
    if reader.is_encrypted:
        raise BuiltinToolError("PASSWORD_REQUIRED", "암호가 설정된 PDF는 처리할 수 없습니다.")
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    # `PdfWriter`는 기본 `/Producer: pypdf`를 넣는다. 빈 dict를 더하는 것으로는
    # 없어지지 않으므로 metadata 자체를 비운다.
    writer.metadata = None
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def _sanitize_docx(data: bytes) -> bytes:
    validate_ooxml_package(data)
    try:
        document = Document(BytesIO(data))
    except (ValueError, KeyError, BadZipFile, InvalidFileException) as exc:
        raise BuiltinToolError("INVALID_DOCX", "손상되었거나 지원하지 않는 DOCX입니다.") from exc
    output = BytesIO()
    document.save(output)
    return _scrub_ooxml_properties(output.getvalue())


def _sanitize_xlsx(data: bytes) -> bytes:
    validate_ooxml_package(data)
    try:
        workbook = load_workbook(BytesIO(data), data_only=False)
    except (ValueError, KeyError, BadZipFile) as exc:
        raise BuiltinToolError("INVALID_XLSX", "손상되었거나 지원하지 않는 XLSX입니다.") from exc
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return _scrub_ooxml_properties(output.getvalue())


_EMPTY_CORE_PROPERTIES = (
    b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    b'<cp:coreProperties '
    b'xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
    b'xmlns:dc="http://purl.org/dc/elements/1.1/">'
    b'<dc:title/><dc:subject/><dc:creator/><cp:keywords/><dc:description/>'
    b'<cp:lastModifiedBy/><cp:revision/><cp:category/><cp:contentStatus/>'
    b'</cp:coreProperties>'
)
_EMPTY_EXTENDED_PROPERTIES = (
    b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    b'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
    b'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"/>'
)
_EMPTY_CUSTOM_PROPERTIES = (
    b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    b'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
    b'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"/>'
)


def _scrub_ooxml_properties(data: bytes) -> bytes:
    """DOCX/XLSX의 내용 entry는 보존하고 속성 XML만 빈 문서로 바꾼다."""

    replacements = {
        "docProps/core.xml": _EMPTY_CORE_PROPERTIES,
        "docProps/app.xml": _EMPTY_EXTENDED_PROPERTIES,
        "docProps/custom.xml": _EMPTY_CUSTOM_PROPERTIES,
    }
    output = BytesIO()
    try:
        with ZipFile(BytesIO(data)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
            for info in source.infolist():
                target.writestr(info, replacements.get(info.filename, source.read(info)))
    except BadZipFile as exc:
        raise BuiltinToolError("INVALID_OOXML", "문서 속성을 정리하지 못했습니다.") from exc
    return output.getvalue()


def sanitize_file(*, data: bytes, mime_type: str) -> bytes:
    """메타데이터가 제거된 새 파일 바이트를 반환한다."""

    if not data:
        raise BuiltinToolError("EMPTY_FILE", "빈 파일은 정리할 수 없습니다.")
    if len(data) > MAX_FILE_BYTES:
        raise BuiltinToolError("FILE_TOO_LARGE", "파일은 20MB 이하여야 합니다.")
    handlers = {PDF_MIME: _sanitize_pdf, DOCX_MIME: _sanitize_docx, XLSX_MIME: _sanitize_xlsx}
    try:
        output = handlers[mime_type](data)
    except KeyError as exc:
        raise BuiltinToolError(
            "UNSUPPORTED_FORMAT", "PDF, DOCX, XLSX 파일의 메타데이터만 제거할 수 있습니다."
        ) from exc
    if len(output) > MAX_OUTPUT_BYTES:
        raise BuiltinToolError("OUTPUT_TOO_LARGE", "결과 파일이 허용 크기를 초과했습니다.")
    if sensitive_metadata(data=output, mime_type=mime_type):
        raise BuiltinToolError(
            "METADATA_REMAINS", "일부 문서 속성을 안전하게 제거하지 못했습니다."
        )
    return output
